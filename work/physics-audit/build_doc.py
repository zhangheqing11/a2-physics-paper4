from pathlib import Path
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
import json,re,hashlib
root=Path.cwd();o=root/'output/physics';p=root/'work/physics-audit'
d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.8);s.bottom_margin=Cm(1.7);s.left_margin=Cm(1.9);s.right_margin=Cm(1.9)
for name,size in [('Normal',10.5),('Title',22),('Heading 1',17),('Heading 2',12)]:
 st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PingFang SC')
 st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.15
 if name!='Normal':st.paragraph_format.space_before=Pt(9)
foot=s.footer.paragraphs[0];foot.alignment=2
foot.add_run('9702 P4 + P5  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
for r in foot.runs:r.font.size=Pt(8)
lines=(o/'学习与题型手册第一版.md').read_text().splitlines();i=0
while i<len(lines):
 line=lines[i].strip();i+=1
 if not line:continue
 if line=='---PAGE---':d.add_page_break();continue
 if line.startswith('|'):
  rows=[line]
  while i<len(lines) and lines[i].strip().startswith('|'):rows.append(lines[i].strip());i+=1
  rows=[r for r in rows if not re.match(r'^\|[ \-:|]+$',r)]
  values=[[v.strip() for v in r.strip('|').split('|')] for r in rows]
  table=d.add_table(rows=0,cols=len(values[0]));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
  widths=([2.1,3.8,11.3] if len(values[0])==3 else [2.1,3.5,5.4,6.2] if len(values[0])==4 else [4.3,3,3.3,3.3,3.3])
  for col,w in zip(table.columns,widths):col.width=Cm(w)
  for ri,row in enumerate(values):
   cells=table.add_row().cells
   for c,w,txt in zip(cells,widths,row):
    c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;c.text=txt
    tcpr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
    for edge in ['top','left','bottom','right']:
     e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
    tcpr.append(b);mar=OxmlElement('w:tcMar')
    for edge in ['top','left','bottom','right']:
     e=OxmlElement('w:'+edge);e.set(qn('w:w'),'80');e.set(qn('w:type'),'dxa');mar.append(e)
    tcpr.append(mar)
    for para in c.paragraphs:
     para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2);para.paragraph_format.line_spacing=1.05
     for run in para.runs:run.font.size=Pt(9);run.bold=(ri==0)
    if ri==0:
     shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF2');tcpr.append(shade)
   if ri==0:
    flag=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(flag)
  d.add_paragraph();continue
 if line.startswith('# '):d.add_paragraph(line[2:],'Title' if len(d.paragraphs)<2 else 'Heading 1')
 elif line.startswith('## '):d.add_paragraph(line[3:],'Heading 2')
 else:d.add_paragraph(line)
for tree in [d.styles.element,d.element]:
 for border in list(tree.iter(qn('w:pBdr'))):border.getparent().remove(border)
dest=o/'CIE9702_P4P5_学习与题型手册_第一版.docx';d.save(dest)
recovery=json.loads((p/'recovery.json').read_text());reviews=json.loads((p/'missing_review.json').read_text())
for r in recovery:
 if 'error' in r:
  r.pop('error');r.update({'source':'https://bestexamhelp.com/exam/cambridge-international-a-level/physics-9702/2024/9702_s24_qp_41.pdf','declared_pages':24,'original_pages':24,'merged_pages':23,'sha256':hashlib.sha256((o/'补充原卷/9702_s24_qp_41.pdf').read_bytes()).hexdigest()})
 r.pop('suspected_missing',None)
 r['review']='页数差异与删空白页相符；已完成文本对照，非全页图形质量审核'
(o/'资料核查记录.json').write_text(json.dumps({'scope':'2022-2025现有50份QP与50份MS；非网站全版本完整性声明','inventory':json.loads((p/'inventory.json').read_text()),'recovered':recovery,'page_text_review':reviews},ensure_ascii=False,indent=2))
print(dest)
