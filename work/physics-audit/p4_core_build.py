from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE=Path(__file__).resolve().parent
OUT=BASE.parents[1]/'output/physics/p4-complete/CIE9702_P4_核心概念与公式总表.docx'

def el(tag, *children):
    e=OxmlElement('m:'+tag)
    for c in children:e.append(c)
    return e

def mr(text):
    r=el('r'); pr=OxmlElement('w:rPr')
    font=OxmlElement('w:rFonts'); font.set(qn('w:ascii'),'Cambria Math');font.set(qn('w:hAnsi'),'Cambria Math');pr.append(font)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'22');pr.append(sz);r.append(pr)
    t=el('t');t.text=text;t.set(qn('xml:space'),'preserve');r.append(t);return r

class Math:
    def __init__(self,s):self.s=s;self.i=0
    def group(self):
        assert self.s[self.i]=='{',(self.s,self.i)
        self.i+=1;out=self.parse('}');self.i+=1;return out
    def atom(self):
        if self.s[self.i]=='{':return self.group()
        if self.s[self.i]=='\\':
            self.i+=1;m=re.match('[A-Za-z]+',self.s[self.i:]);assert m,self.s
            name=m.group();self.i+=len(name)
            if name=='frac':return [el('f',el('num',*self.group()),el('den',*self.group()))]
            if name=='sqrt':
                pr=el('radPr');hide=el('degHide');hide.set(qn('m:val'),'1');pr.append(hide)
                return [el('rad',pr,el('deg'),el('e',*self.group()))]
            if name in ('overline','dot'):
                if name=='overline':
                    pr=el('barPr');pos=el('pos');pos.set(qn('m:val'),'top');pr.append(pos)
                    return [el('bar',pr,el('e',*self.group()))]
                pr=el('accPr');ch=el('chr');ch.set(qn('m:val'),'̇');pr.append(ch)
                return [el('acc',pr,el('e',*self.group()))]
            return [mr({'langle':'〈','rangle':'〉','approx':'≈','gg':'≫'}[name])]
        ch=self.s[self.i];self.i+=1;return [mr(ch)]
    def parse(self,end=None):
        out=[]
        while self.i<len(self.s) and self.s[self.i]!=end:
            a=self.atom();sub=sup=None
            while self.i<len(self.s) and self.s[self.i] in '_^':
                op=self.s[self.i];self.i+=1;b=self.atom()
                if op=='_':sub=b
                else:sup=b
            if sub is not None and sup is not None:a=[el('sSubSup',el('e',*a),el('sub',*sub),el('sup',*sup))]
            elif sub is not None:a=[el('sSub',el('e',*a),el('sub',*sub))]
            elif sup is not None:a=[el('sSup',el('e',*a),el('sup',*sup))]
            out.extend(a)
        return out

d=Document();s=d.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11)
s.top_margin=Inches(.7);s.bottom_margin=Inches(.7);s.left_margin=Inches(.8);s.right_margin=Inches(.8)
s.footer_distance=Inches(.3)
for name,size in [('Normal',11),('Title',24),('Subtitle',12),('Heading 1',17),('Heading 2',12)]:
    st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PingFang SC')
    p=st.paragraph_format;p.space_after=Pt(6);p.line_spacing=1.08
    if name.startswith('Heading'):p.keep_with_next=True;p.space_before=Pt(13)
for st in d.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
st=d.styles.add_style('Equation',1);st.base_style=d.styles['Normal'];st.paragraph_format.space_before=Pt(5);st.paragraph_format.space_after=Pt(4);st.paragraph_format.keep_with_next=True
st=d.styles.add_style('Condition',1);st.base_style=d.styles['Normal'];st.font.size=Pt(10.5);st.paragraph_format.space_after=Pt(8)
st=d.styles.add_style('Source',1);st.base_style=d.styles['Normal'];st.font.size=Pt(9);st.paragraph_format.space_after=Pt(7)
f=s.footer.paragraphs[0];f.alignment=2
r=f.add_run('9702 P4 核心概念与公式   ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
content=(BASE/'p4_core.md').read_text();equations=0;h1=[]
for i,line in enumerate(content.splitlines()):
    line=line.strip()
    if not line:continue
    if line.startswith('# '):
        text=line[2:]
        p=d.add_paragraph(text,'Title' if i==0 else 'Heading 1')
        if i>0:
            h1.append(text)
            p.paragraph_format.page_break_before=(text=='通用工具与 AS 接口')
    elif line.startswith('## '):d.add_paragraph(line[3:],'Heading 2')
    elif line.startswith('@'):
        # Absolute-value bars occur inside formulas; split only at the labelled ends.
        first=line.index('|');last=line.rindex('|')
        label=line[1:first];formula=line[first+1:last];note=line[last+1:]
        assert label in ['核心','推导','接口']
        p=d.add_paragraph(style='Equation');r=p.add_run(label+'  ');r.font.size=Pt(9);r.bold=True
        p._p.append(el('oMath',*Math(formula).parse()));equations+=1
        d.add_paragraph(note,'Condition')
    elif line.startswith('对应考纲'):d.add_paragraph(line,'Source')
    elif i==1:d.add_paragraph(line,'Subtitle')
    else:d.add_paragraph(line)
d.core_properties.title='CIE 9702 P4 核心概念与公式'
d.core_properties.subject='2025至2027考纲第12至25章 学习与刷题查漏'
d.core_properties.author=''
OUT.parent.mkdir(parents=True,exist_ok=True);d.save(OUT)
print({'output':str(OUT),'equation_blocks':equations,'chapters':h1,'paragraphs':len(d.paragraphs)})
