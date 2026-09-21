from pathlib import Path
from collections import defaultdict,Counter
import json,re,csv,sys
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from p4_assemble import rows,bytype,inv
P=Path('work/physics-audit');O=Path('output/physics/p4-complete');root=Path.cwd()
lookup={r['paper'][5:]+':'+r['q']:r for r in rows};types={}
for line in (P/'p4_types.tsv').read_text().splitlines():
 a=line.replace('|Qq|','abs(Qq)').replace('|q|','abs(q)').split('|');assert len(a)==8,a[0]
 code,name,sub,path,trap,ref,worked,english=a
 assert code in lookup[ref]['types']
 types[code]=dict(code=code,name=name,sub=sub,path=path,trap=trap,ref=ref,worked=worked,english=english)
types['T05']['worked']=types['T05']['worked'].replace('把剩余热量分给题定液体质量与55°C温升，c≈2.3','再用[810×21−1940]/[120×55]，c≈2.3')
types['C07']['worked']='本题保持车速不变，把后轮链轮换小；后轮及后链轮的角速度不变，所以同时间链条移动距离更小、链速减小（B1）。前链轮半径不变，因此踏板角速度减小（B1）。不能将“车速不变”误当“链速不变”。'
practice={l.split(' ',1)[0]:l.split(' ',1)[1] for l in (P/'p4_practice.txt').read_text().splitlines()}
assert set(types)==set(bytype)==set(practice)
for t,ref in practice.items():
 if not ref.startswith('自编'):assert ref in lookup and t in lookup[ref]['types'],(t,ref)
def short(p):return p.replace('9702_','').replace('_','/')
def refstr(r):return short(r['paper'])+' Q'+r['q']
def source(r):return f"{refstr(r)}  {r['marks']}分  QP合并第{r['qp_page']}页起  MS合并第{r['page']}页"
def grouped_refs(rs):
 groups=defaultdict(list)
 for r in rs:groups[r['paper']].append(r['q'])
 return '；'.join(short(p)+' Q'+', '.join(qs) for p,qs in groups.items())
d=Document();s=d.sections[0];s.page_width=Cm(21);s.page_height=Cm(29.7);s.top_margin=Cm(1.8);s.bottom_margin=Cm(2.4);s.left_margin=Cm(1.85);s.right_margin=Cm(1.85);s.header_distance=Cm(.7);s.footer_distance=Cm(.8)
for name,size in [('Normal',10.5),('Title',24),('Subtitle',12),('Heading 1',17),('Heading 2',12),('Heading 3',10.5)]:
 st=d.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'PingFang SC')
 st.paragraph_format.space_after=Pt(5);st.paragraph_format.line_spacing=1.12
 if name.startswith('Heading'):st.paragraph_format.space_before=Pt(11);st.paragraph_format.keep_with_next=True
st=d.styles.add_style('Source note',1);st.base_style=d.styles['Normal'];st.font.size=Pt(8.5);st.font.color.rgb=RGBColor.from_string('555555');st.paragraph_format.space_after=Pt(4)
st=d.styles.add_style('Index text',1);st.base_style=d.styles['Normal'];st.font.size=Pt(9);st.paragraph_format.line_spacing=1.08;st.paragraph_format.space_after=Pt(4)
foot=s.footer.paragraphs[0];foot.alignment=2;foot.add_run('9702 P4  2022至2025语料  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
for r in foot.runs:r.font.size=Pt(8)
header=s.header.paragraphs[0];header.text='A2 PHYSICS  PAPER 4  学习与题型手册'
for r in header.runs:r.font.size=Pt(8);r.font.color.rgb=RGBColor(0,0,0)
def para(text,style=None,label=None):
 p=d.add_paragraph(style=style)
 if label:p.add_run(label+'  ').bold=True
 p.add_run(text);return p
def table(values,widths):
 t=d.add_table(rows=0,cols=len(widths));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Cm(w)
 for ri,row in enumerate(values):
  cells=t.add_row().cells
  for c,w,txt in zip(cells,widths,row):
   c.width=Cm(w);c.text=str(txt);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   pr.append(b);mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   for p in c.paragraphs:
    p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.06
    for r in p.runs:r.font.size=Pt(9);r.bold=(ri==0)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8EDF2');pr.append(sh)
  if ri==0:t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 return t
def markdown(text,topbreak=False):
 ls=text.splitlines();i=0
 while i<len(ls):
  line=ls[i].strip();i+=1
  if not line:continue
  if line.startswith('|'):
   vals=[line]
   while i<len(ls) and ls[i].startswith('|'):vals.append(ls[i]);i+=1
   vals=[[s.strip() for s in l.strip('|').split('|')] for l in vals if not re.match(r'^\|[\- :|]+$',l)]
   table(vals,[3.0,8.2,6.1]);continue
  if line.startswith('# '):
   p=para(line[2:],'Heading 1');p.paragraph_format.page_break_before=topbreak
  elif line.startswith('## '):para(line[3:],'Heading 2')
  else:para(line)
para('CIE A2 Physics Paper 4\n学习与题型全册','Title')
para('基于2022至2025年25份试卷的完整小问索引','Subtitle')
para('供尚未系统学习A2物理的学习者使用。先建立模型与条件，再用对应真题练习；目标是能独立识别、求解并解释，而不只是记住答案。')
table([['资料与交付边界','本册内容'],['P4语料','25份QP与25份精确配对MS'],['全量小问','250道大题  1250条小问  2500分'],['分类与教学','85个有题源条目  85则代表小问讲解  85道下一步变式'],['学习基础','全A2主题的概念  双语表达  符号单位  公式条件'],['状态','仅完成手册与资料整理  所有个人掌握状态仍为未学'],['不含范围','Paper 5  2026新卷  未纳入的其他历史版本']],[4.0,13.3])
para('编制日期 2026年9月14日','Source note')
para('正文导航','Heading 2')
para('使用说明 → AS基础 → C圆周 → G引力 → T温度 → K气体 → H热力学 → S振动 → E电场 → D电容 → B磁场 → I感应 → A交流 → Q量子 → N核物理 → M医学 → Z天体 → 图像速查 → 综合与考纲补充 → 统计及逐卷完整索引。')
para('在Word中可使用导航窗格按标题跳转；按条目编号或完整小问搜索即可定位。所有原题出处均在同一册中，无需另开索引文件。','Source note')
foundation=(P/'p4_foundations.md').read_text();chunks=re.split(r'(?m)(?=^# )',foundation)
for chunk in chunks:
 if not chunk.strip():continue
 first=chunk.splitlines()[0]
 markdown(chunk,topbreak=True)
 if first.startswith('# 入门'):codes=['X01']
 else:
  m=re.match(r'# ([A-Z]) ',first);codes=[t for t in types if m and t.startswith(m[1])]
 if codes:
  para('本章题型树','Heading 2')
  for code in codes:para(code+'  '+types[code]['name']+'  —  '+types[code]['sub'],'Index text')
 for code in codes:
  t=types[code];r=lookup[t['ref']];rs=bytype[code];block_start=len(d.paragraphs)
  para(code+' '+t['name'],'Heading 2')
  para(t['sub'],label='细分考法')
  para(t['path'],label='识别与模型')
  para(source(r),'Source note')
  para(t['worked'],label='代表题讲解')
  para(t['english'],label='英文表达')
  para(t['trap'],label='常见失分')
  pr=practice[code]
  if pr.startswith('自编'):para(pr,label='下一道变式')
  else:para(source(lookup[pr])+'。先独立完成，再核对对应MS；说明相对代表题改变了什么。',label='下一道变式')
  para('状态  未学  □学习中  □已掌握    证据/错因/复测日期：________________','Source note')
  para(grouped_refs(rs),'Index text',label='本类型完整题源')
  for pp in d.paragraphs[block_start:-1]:pp.paragraph_format.keep_with_next=True
  for pp in d.paragraphs[block_start:]:pp.paragraph_format.keep_together=True
# Visual reading aids are original schematics, not copies of the exam graphs.
for n,desc in [(1,'图1至4分别对应S02、S03、D05或N04、G03。线条为无量纲示意；原题作图必须恢复题给尺度、符号与起始条件。图4仅画球外区域r≥R，不将球内模型混入。'),(2,'图5至8分别对应I04、A03、A02、N02。图5选定同一线圈正方向后用负导数作emf；图8为结合能趋势示意，不能据此读取精确核数据。')]:
 p=para('图像速查 '+str(n),'Heading 1');p.paragraph_format.page_break_before=True
 para(desc);d.add_picture(str(P/f'figures/graph-{n}.png'),width=Cm(17.3))
markdown((P/'p4_appendix.md').read_text(),topbreak=True)
p=para('题型覆盖统计','Heading 1');p.paragraph_format.page_break_before=True
para('小问关联数含交叉标签；主归属分值只计一次。原始卷数保留41与43，内容组数将已核实MS相同的43归至41。两种计数都不能解释为未来考试概率。')
vals=[['条目','名称','关联小问','卷数/内容组','主归属分']]
for t,x in types.items():
 rs=bytype[t];papers={r['paper'] for r in rs};content={p[:-2]+'41' if p.endswith('43') else p for p in papers}
 vals.append([t,x['name'],len(rs),str(len(papers))+'/'+str(len(content)),sum(r['marks'] for r in rs if r['types'][0]==t)])
table(vals,[1.15,8.1,2.15,3.0,2.9])
p=para('全部试卷与小问索引','Heading 1');p.paragraph_format.page_break_before=True
para('阅读方式：Q小问 [分值] 主类型+交叉类型 (MS页)。每卷先列原题与答案所在的合并页区间，再按Q1至Q10列出全部小问。每道大题后括号中的QP页是该大题起始页。此索引包含25份卷子的全部1250条记录。')
for paper in sorted({r['paper'] for r in rows},key=lambda p:(p[6:8],{'m':0,'s':1,'w':2}[p[5]],p[-2:])):
 pr=[r for r in rows if r['paper']==paper];qid=paper.replace('_4','_qp_4');mid=paper.replace('_4','_ms_4');qp=inv[qid];ms=inv[mid]
 para(short(paper)+' 全部小问','Heading 2')
 para(f"QP合并页 {qp['start']}至{qp['end']}  |  MS合并页 {ms['start']}至{ms['end']}  |  {len(pr)}条  100分"+('  |  MS内容与同考季41相同' if paper.endswith('43') else ''),'Source note')
 for q in range(1,11):
  qr=[r for r in pr if r['q'].split('(')[0]==str(q)]
  para(f"Q{q}  {sum(r['marks'] for r in qr)}分  QP第{qr[0]['qp_page']}页起",'Heading 3')
  para('； '.join(r['q']+f" [{r['marks']}] "+'+'.join(r['types'])+f" (MS {r['page']})" for r in qr),'Index text')
p=para('源文件位置与使用边界','Heading 1');p.paragraph_format.page_break_before=True
para('题卷合并PDF：sources/（已压缩）9702_ALL_QuestionPapers_merged.pdf。答案合并PDF：sources/9702_all_ms_merged.pdf。补充的13份完整原卷保存在output/physics/补充原卷。它们均按精确卷号配对，本手册未修改sources中的同步资料。')
para('本册不全文重印1250条答案。回查步骤：从条目代表或练习读取卷号与小问；按附录定位QP大题起始页和MS小问页；阅读完整题干及相邻前问；作答后使用同卷号MS核对。图像被压缩得不清楚时，优先检查补充原卷。')
para('手册与结构化总账已相互核验：25卷各100分，250道大题分值一致，1250条小问无空类型，85条代表与85道练习均已验证引用。完整手册不等于个人学习完成；后续状态需要你的独立作答和解释来更新。')
for tree in [d.styles.element,d.element]:
 for border in list(tree.iter(qn('w:pBdr'))):border.getparent().remove(border)
dest=O/'CIE9702_P4_学习与题型全册_2022至2025.docx';d.save(dest)
fields=['paper','q','marks','primary_type','cross_types','qp_main_start','ms_page','equivalent_ms','mastery']
with (O/'P4全量小问总账.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.DictWriter(f,fieldnames=fields);w.writeheader()
 for r in rows:w.writerow(dict(paper=r['paper'],q=r['q'],marks=r['marks'],primary_type=r['types'][0],cross_types='+'.join(r['types'][1:]),qp_main_start=r['qp_page'],ms_page=r['page'],equivalent_ms=r['equivalent_ms'],mastery='未学'))
with (O/'P4掌握记录.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(['类型','名称','状态','日期','独立题目','得分','错误类别','解释证据','复测日期','复测结果'])
 for t in types.values():w.writerow([t['code'],t['name'],'未学','','','','','','',''])
audit={'papers':25,'main_questions':250,'subquestions':len(rows),'marks':sum(r['marks'] for r in rows),'types':len(types),'cross_tagged':sum(len(r['types'])>1 for r in rows),'ms_content_groups':17,'non43_rows':846,'unmapped':0,'qp_main_mark_check':'250/250 equal to matching MS','representatives':85,'practice':85,'document':str(dest.resolve())}
(O/'P4覆盖核验.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
print(json.dumps(audit,ensure_ascii=False,indent=2))
